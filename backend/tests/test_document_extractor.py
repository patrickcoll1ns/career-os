from io import BytesIO

from docx import Document as DocxDocument

from app.integrations.document_extractor import DocumentTextExtractor

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def test_extracts_and_normalizes_plain_text() -> None:
    data = "  Backend Engineer  \n\n Python and FastAPI \n".encode()

    text = DocumentTextExtractor().extract(data, "text/plain")

    assert text == "Backend Engineer\nPython and FastAPI"


def test_extracts_paragraphs_and_tables_from_docx() -> None:
    document = DocxDocument()
    document.add_paragraph("Backend Engineer")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "FastAPI"
    buffer = BytesIO()
    document.save(buffer)

    text = DocumentTextExtractor().extract(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert text == "Backend Engineer\nPython | FastAPI"


def test_strips_a_utf8_byte_order_mark() -> None:
    text = DocumentTextExtractor().extract("﻿Backend Engineer".encode(), "text/plain")

    assert text == "Backend Engineer"
